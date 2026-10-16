"""Create the updated PBL report with the retained university cover pages."""
import copy
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/multisensor_v2'


def main():
    summary=json.loads((OUT/'test_summary.json').read_text())
    evidence=json.loads((OUT/'evidence_summary.json').read_text())
    primary=summary['primary_test'];frozen=json.loads((OUT/'frozen_model.json').read_text())
    selector=json.loads((ROOT/'results/multisensor_v1/feature_selection.json').read_text())
    doc=Document(ROOT/'submission/HVAC_PBL_Project_Report.docx')
    body=doc._element.body;children=list(body);section=copy.deepcopy(body.sectPr)
    for element in children[69:]:body.remove(element)
    body.append(section)
    for level,size in [(1,16),(2,12)]:
        style=doc.styles.add_style(f'Multi Heading {level}',WD_STYLE_TYPE.PARAGRAPH)
        style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=True
        style.paragraph_format.space_after=Pt(10);style.paragraph_format.keep_with_next=True
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1));style.element.get_or_add_pPr().append(outline)
    for name in ['Multi Heading 1','Multi Heading 2','PBL Body','PBL Table','PBL Caption']:
        style=doc.styles[name];style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['PBL Body'].font.size=Pt(11.5)
    doc.styles['PBL Body'].paragraph_format.line_spacing=1.12
    doc.styles['PBL Body'].paragraph_format.space_after=Pt(7)
    doc.styles['PBL Table'].font.size=Pt(10)
    pct=lambda v:f'{v:.2%}'

    def para(text):return doc.add_paragraph(text,'PBL Body')
    def page(title):
        p=doc.add_paragraph(title,'Multi Heading 1');p.paragraph_format.page_break_before=True
        return p
    def sub(title):return doc.add_paragraph(title,'Multi Heading 2')
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for i,values in enumerate([headers]+rows):
            row=t.rows[0] if i==0 else t.add_row()
            for j,text in enumerate(values):
                cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.style=doc.styles['PBL Table'];r=p.add_run(str(text));r.bold=i==0
                if j>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                pr=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF2' if i==0 else ('F7F9FA' if i%2==0 else 'FFFFFF'));pr.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'90');edge.set(qn('w:type'),'dxa');margins.append(edge)
                pr.append(margins)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    edge=OxmlElement('w:'+side)
                    for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:edge.set(qn('w:'+key),value)
                    borders.append(edge)
                pr.append(borders)
            if i==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph('',style='PBL Caption')
        return t
    def figure(name,caption,width=6.1):
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(OUT/name),width=Inches(width))
        doc.add_paragraph(caption,'PBL Caption')

    page('Abstract')
    para('This project extends deep learning fault detection to four sensor locations in a simulated cooling and heating air-handling unit. We train an LSTM to distinguish healthy operation, cold-deck air temperature sensor bias, hot-deck air temperature sensor bias, cold-deck duct pressure sensor bias and hot-deck duct pressure sensor bias. Each scenario contains one faulty sensor. The earlier outdoor-air detector remains a separate experiment in the same repository.')
    para('The LBNL dual-duct AHU source provides a full year of one-minute measurements. We use seventeen native scenarios and twenty-six recorded inputs. Whole calendar weeks stay together across every parallel scenario, which prevents the same weather period from entering more than one partition. Cleaning and scaling use healthy training rows. We also evaluate five inputs selected without validation or test labels.')
    para(f"The primary model achieves {pct(primary['accuracy'])} five-class test accuracy and {pct(primary['macro_f1'])} macro F1 across {summary['test_windows']:,} scored minutes. Correct-sensor recall ranges from {pct(min(r['recall'] for r in primary['per_class'][1:]))} to {pct(max(r['recall'] for r in primary['per_class'][1:]))}. Healthy FPR is {pct(primary['healthy_fpr'])}, exceeding the 1% goal. The five-input alternatives perform worse and do not replace the full-input model. The test remains fixed after inspection.")
    para('The submission includes saved model weights, a Streamlit demonstration, a reproducible notebook, preprocessing definitions, source and partition audits, complete score arrays and verification scripts. This is an academic simulation study. It does not establish simultaneous-fault diagnosis, reliability on physical equipment or future-failure forecasting.')
    sub('Project details')
    para('DSE 3170 Project Based Learning 3, Section B. Team: Rishabh Srivastava 2430010149, Mitul Seth 2430010099 and Madhav Mittal 2430010096. Guide: Dr. Neha V Sharma, Assistant Professor (Selection Grade). Class instructor: Mr. Santosh Rajak, Assistant Professor. October 2026, academic year 2026 to 2027.')
    para('Keywords: deep learning, LSTM, sensor fault diagnosis, HVAC simulation, feature selection, false alarms.')

    page('1 Introduction and objectives')
    para('An air-handling unit distributes conditioned air through a building. A dual-duct system uses a cold deck and a hot deck, with terminal units mixing air to meet room requirements. Incorrect sensor readings can cause controls to supply the wrong amount of heating, cooling or airflow. A useful detector should identify the affected sensor while keeping alarms on healthy operation low.')
    para('The earlier project accepted several sensor measurements but learned only outdoor-air temperature bias. This extension changes the diagnostic targets. It studies four different sensor locations using native labeled cases from a different simulated AHU. Adding inputs and adding fault targets are separate design choices.')
    table(['Output','Meaning'],[['Healthy','No modeled sensor bias'],['Cold-deck temperature sensor','Bias in the cold supply-air temperature sensor'],['Hot-deck temperature sensor','Bias in the hot supply-air temperature sensor'],['Cold-deck pressure sensor','Bias in the cold supply-duct pressure sensor'],['Hot-deck pressure sensor','Bias in the hot supply-duct pressure sensor']],[2.5,3.8])
    sub('Objectives')
    para('We aim to identify each of the four fault types, compare a wider measurement set with five selected raw inputs, and measure false positives on healthy operation. Validation targets require healthy FPR at or below 1% and correct-sensor recall of at least 40% for every fault type. Five-class accuracy and macro F1 summarize the classification task, while sensor-specific precision and recall reveal unequal performance.')
    para('The scope covers one simulated unit, known single-sensor bias families and held-out calendar periods. It does not retain outdoor-air bias as one of the four new classes. The dashboard keeps the earlier outdoor-air benchmark separately available.')

    page('2 Dataset and source audit')
    para('We use the official LBNL dual-duct AHU collection, generated with HVACSIM+. The downloaded archive contains 56 CSV files and 114 recorded measurement or control points. We select one fault-free file and sixteen sensor-bias files, covering four native severity identifiers per sensor location. Each selected file contains 525,600 one-minute rows in 2018. These are parallel scenarios of one modeled system, rather than seventeen independent units.')
    table(['Property','Study value'],[['Selected source rows','8,935,200'],['Scenario files','1 healthy and 16 sensor-bias cases'],['Model classes','5'],['Model inputs','26 recorded AHU channels'],['Input measurements and commands','21 measurements and 5 commands'],['History and sampling','30 readings at one-minute cadence']],[3.4,2.9])
    para('Temperature values use Fahrenheit. The paper documents temperature offsets of plus or minus 2 and 4 Celsius. Pressure severity filenames say plus or minus 2inwg and 4inwg, while Table 4 of the source paper describes plus or minus 0.2 and 0.4 inH2O. We retain native scenario identifiers and do not assert independently verified pressure offset magnitudes.')
    sub('Rejected candidate source')
    para('An initial single-duct AHU download failed integrity checks. Its four outdoor-air bias tables were identical, and the documented supply-air bias filenames were absent. We retained that audit and excluded those cases from this study. The chosen dual-duct sensor cases have distinct archive signatures and separate trace data.')
    para('Source: LBNL Fault Detection and Diagnostics Datasets, collection DOI 10.25984/1881324. Dataset page: https://faultdetection.lbl.gov/dataset/simulated-dd-ahu-dataset/. Full provenance and archive hashes accompany the repository.')

    page('3 Experimental methodology')
    para('We assign entire ISO calendar weeks using a fixed rule. Weeks whose number modulo six equals zero form the test partition; a remainder of five forms validation; all other weeks form training. Every source scenario uses the same calendar assignment. This keeps matching weather and operating periods together and avoids randomly splitting overlapping histories.')
    table(['Partition','Eligible endpoints','Used for optimization'],[['Training',f"{evidence['eligible_train_endpoints']:,}",f"{evidence['sampled_train_windows']:,} hourly endpoints"],['Validation',f"{evidence['eligible_validation_endpoints']:,}",f"{evidence['checkpoint_validation_windows']:,} quarter-hour checkpoints"],['Test',f"{evidence['eligible_test_endpoints']:,}",'Every eligible minute after freezing']],[1.25,1.8,3.25])
    para('Policy selection also uses every eligible validation minute. Hourly training endpoints reduce duplicated optimization work while retaining the original one-minute history within each input window. The larger eligible-row count must not be described as the number of independently optimized training examples.')
    sub('Operating eligibility and cleaning')
    para('A valid window contains thirty contiguous one-minute readings in one source day, with SYS_CTL equal to one and both hot and cold supply fan speeds above zero. Histories reset at scenario, day and cadence boundaries. Shutdown periods and other inactive readings remain unscored. Scenario-level raw and scored day counts preserve this coverage limitation.')
    para('We preserve valid negative outdoor Fahrenheit readings. Nonfinite values receive causal forward filling for at most five minutes within a continuous source day. Remaining gaps use median values fitted to healthy active training rows. Standardization uses the same training rows, followed by clipping at ten standard deviations. Timestamp, scenario identity, filename severity and class labels never enter the model.')
    para('Evaluation distinguishes identifying the correct sensor from merely raising any fault alarm. A wrong sensor alarm counts as a classification error even when it detects an abnormal condition. Healthy FPR uses only healthy minutes in its denominator.')

    page('4 Feature selection and deep learning model')
    para('The wider input set contains five air temperatures, two duct pressures, four humidity measurements, four airflows, three fan speeds, three fan power measurements and five actuator commands. This gives twenty-one measurement channels and five commands. The model uses all twenty-six raw channels without the derived features from the earlier outdoor-air benchmark.')
    para('A tree selector fits a training-only internal split. Its grouped permutation importance uses held-out training weeks and moves complete daily trajectories together across latest, mean, standard deviation and slope statistics. Three fixed permutation repeats rank the twenty-six inputs. Validation and test data do not determine this ranking. Correlations make the ranking conditional on the selector, rather than a globally optimal subset.')
    table(['Selected raw input','Type'],[['Hot supply fan speed','Measurement'],['Hot supply fan power','Measurement'],['Cold supply fan speed','Measurement'],['Cooling-coil valve command','Control command'],['Cold supply fan power','Measurement']],[4.5,1.8])
    table(['Model component','Configuration'],[['Sequence encoder','Single LSTM layer with 32 hidden units'],['Classification head','16-unit ReLU, dropout 0.2, 5 output logits'],['Primary trainable parameters','8,293'],['Optimizer and batch','Adam 0.001, batch 512, gradient clipping 1'],['Loss','Weighted cross entropy, extra healthy weight in V2'],['Reproducibility','Fixed seed 42 and saved tensor checkpoints']],[2.3,4.0])
    para('Softmax outputs provide class scores. They are uncalibrated and do not establish confidence or deployment probabilities. Five selected raw inputs comprise four measurements and one command; they are not five different fault-target sensors.')

    page('5 Validation and model selection')
    para('Version 1 trained the two input sets for up to twenty epochs and selected checkpoints using fault macro average precision. It failed the combined false-alarm and per-sensor recall validation targets. Before opening the reserved test, Version 2 logged a development refinement: an eighty-epoch maximum, eight-epoch patience, twice the healthy loss weight and checkpoint selection using macro average precision across all five classes.')
    para('Four unique LSTMs were trained in total. Version 2 compares both new models and the two reused Version 1 models. This exposes the development history instead of presenting a favorable final run as the only attempted model. Only seed 42 is tested in this extension; repeat-seed robustness remains future work.')
    para('The final calibration study evaluates 126 validation policies per candidate, or 504 policies across four candidates. Per-sensor thresholds use healthy-score percentiles. Some policies assign a larger false-alarm budget to one sensor while applying stricter thresholds to the others. Immediate, two-consecutive and two-of-three rules share boundary-reset rules. Validation macro F1 selects among policies satisfying the 1% healthy FPR goal and 40% minimum recall for every faulty sensor.')
    para(f"The selected full-input model restores epoch {evidence['best_epoch']} after {evidence['epochs_run']} epochs. Its validation accuracy is {pct(frozen['policy']['validation_metrics']['accuracy'])}, macro F1 is {pct(frozen['policy']['validation_metrics']['macro_f1'])}, and healthy FPR is {pct(frozen['policy']['validation_metrics']['healthy_fpr'])}. Its selected policy uses immediate sensor-specific thresholds.")
    para('Model tensors, preprocessing definitions, thresholds, feature choice and primary candidate were frozen before final test inference. Both feature sets and both development versions remain in the comparison. No post-test threshold changes alter the reported result.')
    table(['Sensor score','Frozen threshold'],[[r['sensor'],f"{t:.6f}"] for r,t in zip(primary['per_class'][1:],frozen['policy']['thresholds'])],[4.4,1.9])

    page('6 Reserved test results')
    para(f"The primary model correctly identifies the five-way source condition in {pct(primary['accuracy'])} of {summary['test_windows']:,} eligible test minutes. Macro F1 is {pct(primary['macro_f1'])}; balanced accuracy is {pct(primary['balanced_accuracy'])}. Any-fault detection recall is {pct(primary['fault_detection_recall'])}, which differs from correct-sensor recall.")
    table(['Faulty sensor','Precision','Recall','F1'],[[r['sensor'].replace(' sensor',''),pct(r['precision']),pct(r['recall']),pct(r['f1'])] for r in primary['per_class'][1:]],[3.0,1.1,1.1,1.1])
    figure('test_sensor_performance.png','Figure 1 Precision and correct-sensor recall for the four fault types',6.1)
    para(f"Healthy false-positive rate is {pct(primary['healthy_fpr'])}, or {primary['healthy_false_readings']:,} false readings among {primary['healthy_readings']:,} healthy readings. The 1% target is not met. The 580 false readings form 52 bursts across forty active healthy calendar days. This is a research result with a remaining false-alarm problem.")

    page('7 Confusion counts and error interpretation')
    figure('test_confusion_matrix.png','Figure 2 Five-class confusion matrix with native counts and within-class percentages',6.1)
    para('Diagonal entries identify the correct condition. Off-diagonal entries in the healthy row are false alarms. Entries in the healthy prediction column for faulted rows are missed faults. Off-diagonal entries between fault classes identify the wrong faulty sensor. These failure types have different operational meanings.')
    para('The hot-deck temperature sensor has the lowest correct-sensor recall, 87.52%. Of the 580 false alarms on healthy minutes, 490 identify this sensor. The healthy-class precision is 47.22% because many faulted minutes receive a healthy prediction in a deliberately fault-heavy experimental sample. Reporting only fault-class precision would hide those misses.')
    para(f"The test fault prevalence is {pct(evidence['test_fault_prevalence'])}. A constant majority-class predictor would score {pct(evidence['five_class_majority_baseline_accuracy'])} on this five-class task. Overlapping minute windows remain correlated. The dataset's constructed prevalence and large number of rows do not demonstrate operational prevalence or independent-equipment reliability.")

    page('8 Comparison with five selected inputs')
    rows=[]
    for family,metric in summary['all_families'].items():
        model=json.loads((OUT/family/'summary.json').read_text())
        label={'v1_all_inputs':'V1 full inputs','v1_top5_inputs':'V1 five inputs','all_inputs':'V2 full inputs','top5_inputs':'V2 five inputs'}[family]
        rows.append([label,str(model['channels']),pct(metric['accuracy']),pct(metric['macro_f1']),pct(metric['healthy_fpr'])])
    table(['Model','Inputs','Accuracy','Macro F1','Healthy FPR'],rows,[1.8,.65,1.2,1.2,1.45])
    para('The full-input Version 2 model was selected using validation before the test. The raw-five alternatives failed validation and perform worse on the reserved weeks. In particular, the selected fan and valve inputs lose information needed for hot-deck temperature identification. Reducing the input count does not automatically improve the false-positive tradeoff.')
    para('The full model uses 8,293 trainable parameters. The five-input architecture is smaller, but this study prioritizes correct sensor identification and false alarms over parameter reduction. A future five-engineered-feature study could retain useful physical relationships while reporting how many physical measurements those features require. The present experiment does not evaluate PCA or every possible five-feature selection.')
    para('The earlier outdoor-air model reported 92.82% accuracy and 1.44% FPR on a different source, task and split. Those values do not form a matched comparison with this four-sensor task. The extension establishes broader modeled fault coverage; it does not demonstrate a reduction in FPR relative to that earlier benchmark.')
    para('A fair follow-up should reserve new data before revising preprocessing, feature selection, architecture or thresholds. The current test weeks are already inspected. We cannot use their error patterns repeatedly and continue describing them as untouched independent evidence.')

    page('9 Application and reproducibility')
    para('The Streamlit dashboard opens the Four-sensor extension by default. Its recorded simulation tab selects a native condition, source scenario and reserved test day. Separate temperature and pressure charts preserve their units. The app compares the model prediction with the author-defined scenario condition and displays the missed false-positive target.')
    para('The upload tab accepts exactly thirty consecutive one-minute readings from one operating unit. The downloadable example supplies the exact columns. Saved training transformations and weights produce class scores; reference labels and filename severity are ignored. The selected immediate policy can produce a current sensor alarm from one complete history. The earlier outdoor-air study remains selectable in the sidebar.')
    sub('Run the demonstration')
    for text in [r'python -m venv .venv',r'.\.venv\Scripts\python.exe -m pip install -r requirements-tested.txt',r'.\.venv\Scripts\python.exe -m streamlit run dashboard\app.py']:
        doc.add_paragraph(text,'PBL Code')
    para('Start Dashboard.bat provides the same demonstration when dependencies are installed. Add --server.port 8502 if port 8501 is already occupied. The saved model, examples and replay run offline; the 1.7 GB dual-duct source archive is downloaded separately for retraining.')
    sub('Reproduce and verify')
    para('Run scripts/download_sda.py --dual, prepare_multisensor.py, train_multisensor.py, refine_multisensor.py, evaluate_multisensor.py and verify_multisensor.py in that order. Each name refers to a Python script inside scripts/. The evaluation reuses saved results on rerun. HVAC_multisensor_submission.ipynb presents the saved evidence without changing the test policy.')
    para('Verification reproduces eight complete validation/test score arrays, all 504 final validation policies and twenty fresh CSV examples with label-exclusion checks. Unit tests cover boundary resets, training-only statistics and the distinction between detection and sensor identification. The source archive and cached tables have recorded hashes.')

    page('10 Conclusion and future scope')
    para('The project now supports fault diagnosis at four temperature and pressure sensor locations in a simulated AHU. The selected full-input LSTM achieves 92.87% five-class accuracy, with correct-sensor recalls between 87.52% and 95.88%. The five selected raw inputs reduce model size but weaken identification. This gives an evidence-based answer to the suggestion that fewer features must improve accuracy.')
    para('The reserved healthy FPR of 2.10% misses the 1% goal despite passing validation. The next model needs a new evaluation design and additional independent evidence. Promising development questions include stronger physical consistency features, further healthy operating examples, calibration across operating modes and a controlled comparison of persistence rules. Those changes must not be selected against the already-inspected test.')
    para('This experiment covers one simulated unit and one known sensor bias at a time. It does not establish simultaneous-fault handling, unseen-severity generalization, physical-equipment reliability, missing-data robustness or future-failure forecasting. Only one training seed is evaluated in the extension. Inactive and shutdown periods remain outside the modeled operating task.')
    sub('References')
    for text in [
        'Granderson et al. A labeled dataset for building HVAC systems operating in faulted and fault-free states. Scientific Data 10, 342 (2023). DOI 10.1038/s41597-023-02197-w. Table 4 defines the four sensor-bias families.',
        'LBNL Fault Detection and Diagnostics Datasets. Dual-duct AHU source and data inventory. https://faultdetection.lbl.gov/dataset/simulated-dd-ahu-dataset/. Collection DOI 10.25984/1881324. Accessed 16 October 2026.',
        'Hochreiter and Schmidhuber. Long Short-Term Memory. Neural Computation 9, 1735 to 1780 (1997). DOI 10.1162/neco.1997.9.8.1735.',
        'Project evidence: results/multisensor_v1 and results/multisensor_v2, including frozen plans, source audits, model tensors, score arrays, comparisons and verification.json. Repository: https://github.com/Mitul-dev/DataCenter-Cooling-Anomaly-Detection.'
    ]:para(text)
    doc.core_properties.title='Cooling Equipment Fault Detection Across Four Sensors'
    doc.core_properties.subject='DSE 3170 PBL 3 multi-sensor extension'
    doc.core_properties.author='Rishabh Srivastava; Mitul Seth; Madhav Mittal'
    output=ROOT/'submission/HVAC_Multisensor_PBL_Report.docx'
    doc.save(output)
    print('Created',output)


if __name__=='__main__':main()
